from datetime import date
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from drugtest_import import extract
from drugtest_import.schema import Extraction, normalize, normalize_date, validate
from drugtest_import.storage import RecordStore

GOOD = {"EmployeeID": "emp005", "Name": " Bob Wilson ", "Department": "Sales", "TestDate": "9/18/2025",
        "TestType": "urine", "Result": "NEGATIVE", "Notes": "Pre-employment screen"}


@pytest.mark.parametrize("raw,expected", [
    ("9/14/2025", "2025-09-14"), ("09-14-25", "2025-09-14"), ("2025-09-14", "2025-09-14"),
    ("Sept 14, 2025", None), ("Sep 14, 2025", "2025-09-14"), ("13/40/2025", None), ("garbage", None),
])
def test_normalize_date(raw, expected):
    assert normalize_date(raw) == expected


def test_normalize_canonicalises_fields():
    record = normalize(GOOD)
    assert record == {"EmployeeID": "EMP005", "Name": "Bob Wilson", "Department": "Sales", "TestDate": "2025-09-18",
                      "TestType": "Urine", "Result": "Negative", "Notes": "Pre-employment screen"}
    assert validate(record, today=date(2025, 10, 1)) == {}


def test_validate_flags_problems():
    record = normalize({**GOOD, "Name": "", "TestDate": "soon", "Result": "Negatve", "TestType": "Sweat"})
    problems = validate(record, today=date(2025, 10, 1))
    assert set(problems) == {"Name", "TestDate", "Result", "TestType"}


def test_validate_rejects_future_date():
    assert "TestDate" in validate(normalize(GOOD), today=date(2025, 9, 1))


@pytest.fixture
def store(tmp_path):
    store = RecordStore(str(tmp_path / "records.db"))
    yield store
    store.close()


def test_store_round_trip_keeps_text_ids(store):
    store.add(normalize({**GOOD, "EmployeeID": "007123"}), source_file="/private/photos/note1.jpg")
    [record] = store.all()
    assert record["EmployeeID"] == "007123"
    assert record["SourceFile"] == "note1.jpg"  # full path not stored
    assert store.count() == 1


def test_find_duplicate(store):
    record = normalize(GOOD)
    assert store.find_duplicate(record) is None
    store.add(record)
    assert store.find_duplicate({**record, "Name": "bob wilson", "EmployeeID": None})["Name"] == "Bob Wilson"
    assert store.find_duplicate({**record, "EmployeeID": "EMP999"}) is None
    assert store.find_duplicate({**record, "TestType": "Hair"}) is None


def test_import_legacy_csv_skips_invalid_rows(store, tmp_path):
    path = tmp_path / "old.csv"
    path.write_text(
        "EmployeeID,Name,Department,TestDate,TestType,Result,Notes\n"
        ",John Doe,HR,9/15/2025,urine,negative,ok\n"
        "EMP9,,IT,2025-09-15,Blood,Pending,missing name\n"
    )
    assert store.import_csv(str(path)) == (1, 1)
    [record] = store.all()
    assert (record["Name"], record["TestDate"], record["TestType"]) == ("John Doe", "2025-09-15", "Urine")


def test_export_xlsx(store, tmp_path):
    store.add(normalize({**GOOD, "EmployeeID": "007123"}))
    store.add(normalize({**GOOD, "Name": "Ann Lee", "Result": "Positive"}))
    path = tmp_path / "out.xlsx"
    assert store.export_xlsx(str(path)) == 2
    ws = load_workbook(path).active
    assert [c.value for c in ws[1]][:7] == ["EmployeeID", "Name", "Department", "TestDate", "TestType", "Result", "Notes"]
    assert ws["A2"].value == "007123"
    assert ws["D2"].value.date() == date(2025, 9, 18)
    assert ws["F3"].value == "Positive"
    assert ws.freeze_panes == "A2"


class FakeClient:
    def __init__(self, response):
        self.calls = []
        self.beta = SimpleNamespace(messages=SimpleNamespace(parse=self._parse))
        self.response = response

    def _parse(self, **kwargs):
        self.calls.append(kwargs)
        return self.response


@pytest.fixture
def image(tmp_path):
    from PIL import Image
    path = tmp_path / "note.png"
    Image.new("RGB", (3000, 2000), "yellow").save(path)
    return str(path)


def test_extract_fields_returns_parsed_output(image):
    parsed = Extraction(Name="Bob Wilson", Result="Negative", uncertain_fields=["Result"])
    client = FakeClient(SimpleNamespace(stop_reason="end_turn", parsed_output=parsed))
    assert extract.extract_fields(image, client=client) is parsed
    content = client.calls[0]["messages"][0]["content"]
    assert content[0]["source"]["media_type"] == "image/jpeg"


@pytest.mark.parametrize("stop_reason", ["refusal", "max_tokens"])
def test_extract_fields_raises_on_bad_stop(image, stop_reason):
    client = FakeClient(SimpleNamespace(stop_reason=stop_reason, parsed_output=None))
    with pytest.raises(extract.ExtractionError):
        extract.extract_fields(image, client=client)
