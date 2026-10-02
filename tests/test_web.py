import io
import re
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from drugtest_import import web
from drugtest_import.crypto import Cipher, generate_key
from drugtest_import.extract import ExtractionError
from drugtest_import.storage import Database

PASSWORD = "a-long-test-password"


def png(color="yellow") -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (60, 40), color).save(buffer, format="PNG")
    return buffer.getvalue()


def fake_extractor(image, form):
    color = Image.open(image).getpixel((0, 0))
    if color == (255, 0, 0):
        raise ExtractionError("blurry")
    if form.key == "equipment-inspection":
        return {"equipment_id": "fl-3", "inspection_date": "3/1/2025", "inspector": "Ana", "status": "needs repair"}, []
    return {"employee_id": "emp005", "name": "Bob Wilson", "test_date": "9/18/2025", "test_type": "urine",
            "result": "Negative"}, ["result"]


@pytest.fixture
def db(tmp_path):
    db = Database(str(tmp_path / "web.db"), Cipher(generate_key()))
    for role in ("admin", "reviewer", "viewer"):
        db.create_user(role, PASSWORD, role, "setup")
    yield db
    db.close()


@pytest.fixture
def app(db, monkeypatch):
    monkeypatch.setenv("DRUGTEST_INSECURE_COOKIES", "1")
    return web.create_app(db, extractor=fake_extractor, workers=1)


def csrf(client, path):
    match = re.search(r'name="csrf" value="([^"]+)"', client.get(path).text)
    return match.group(1)


def login(app, username):
    client = TestClient(app)
    response = client.post("/login", data={"username": username, "password": PASSWORD,
                                           "csrf": csrf(client, "/login")}, follow_redirects=False)
    assert response.status_code == 303
    return client


def wait_until_read(db, timeout=5):
    deadline = time.time() + timeout
    while any(u["status"] == "reading" for u in db.pending_uploads()):
        assert time.time() < deadline, "extraction didn't finish"
        time.sleep(0.02)


def test_pages_require_login(app):
    client = TestClient(app)
    for path in ("/", "/review", "/records", "/admin/users", "/records/export.xlsx"):
        response = client.get(path, follow_redirects=False)
        assert response.status_code == 303 and response.headers["location"] == "/login"


def test_security_headers(app):
    response = TestClient(app).get("/login")
    assert "frame-ancestors 'none'" in response.headers["content-security-policy"]
    assert response.headers["cache-control"] == "no-store"


def test_wrong_password_then_lockout(app, db):
    client = TestClient(app)
    token = csrf(client, "/login")
    for _ in range(web.LOGIN_ATTEMPTS):
        response = client.post("/login", data={"username": "viewer", "password": "nope", "csrf": token})
        assert response.status_code == 401
    response = client.post("/login", data={"username": "viewer", "password": PASSWORD, "csrf": token})
    assert response.status_code == 429
    actions = [e["action"] for e in db.audit_entries()]
    assert actions.count("login.failed") == web.LOGIN_ATTEMPTS and "login.locked" in actions


def test_post_without_csrf_rejected(app):
    client = TestClient(app)
    client.get("/login")
    response = client.post("/login", data={"username": "admin", "password": PASSWORD})
    assert response.status_code == 400


def test_disabled_user_logged_out(app, db):
    client = login(app, "viewer")
    db.update_user(db.get_user_by_name("viewer")["id"], "test", active=False)
    assert client.get("/records", follow_redirects=False).headers["location"] == "/login"


def test_idle_session_expires(app, monkeypatch):
    client = login(app, "viewer")
    assert client.get("/records").status_code == 200
    real_time = time.time
    monkeypatch.setattr(web.time, "time", lambda: real_time() + web.IDLE_TIMEOUT + 1)
    assert client.get("/records", follow_redirects=False).headers["location"] == "/login"


def test_roles(app):
    viewer = login(app, "viewer")
    assert viewer.get("/records").status_code == 200
    assert viewer.get("/review").status_code == 403
    assert viewer.get("/records/export.xlsx").status_code == 403
    assert viewer.get("/admin/users").status_code == 403
    reviewer = login(app, "reviewer")
    assert reviewer.get("/review").status_code == 200
    assert reviewer.get("/admin/audit").status_code == 403


def test_upload_review_save_export(app, db):
    client = login(app, "reviewer")
    token = csrf(client, "/review")
    files = [("files", ("note1.png", png(), "image/png")), ("files", ("bad.png", png("red"), "image/png")),
             ("files", ("notes.txt", b"not an image", "text/plain"))]
    response = client.post("/uploads", data={"csrf": token, "form_key": "drug-test"}, files=files)
    assert response.status_code == 200 and "notes.txt" in response.text  # partial success page
    wait_until_read(db)
    first, second = db.pending_uploads()
    assert second["status"] == "failed" and second["error"] == "blurry"

    page = client.get(f"/review/{first['id']}").text
    assert 'value="2025-09-18"' in page and "Hard to read" in page
    image = client.get(f"/uploads/{first['id']}/image")
    assert image.headers["content-type"] == "image/png" and image.content == png()

    form = {"csrf": token, "action": "save", "field_employee_id": "EMP005", "field_name": "Bob Wilson",
            "field_test_date": "2025-09-18", "field_test_type": "Urine", "field_result": "Negative", "field_notes": ""}
    assert client.post(f"/review/{first['id']}", data={**form, "field_result": "Negatve"}).status_code == 422
    response = client.post(f"/review/{first['id']}", data=form)
    assert response.status_code == 422 and "Tick the box" in response.text
    response = client.post(f"/review/{first['id']}", data={**form, "confirm_uncertain": "1"}, follow_redirects=False)
    assert response.status_code == 303
    assert db.count() == 1 and db.get_upload(first["id"]) is None

    # The failed one: entered by hand, then caught as a duplicate
    response = client.post(f"/review/{second['id']}", data=form)
    assert response.status_code == 409 and "already recorded" in response.text
    client.post(f"/review/{second['id']}", data={**form, "action": "skip"})
    assert db.pending_uploads() == [] and db.count() == 1

    records = client.get("/records?form=drug-test&q=bob").text
    assert "Bob Wilson" in records and "reviewer" in records
    assert "Bob Wilson" not in client.get("/records?form=drug-test&flag=Positive").text
    assert "Bob Wilson" not in client.get("/records?form=drug-test&flag=*flagged").text
    assert "Bob Wilson" not in client.get("/records?form=drug-test&start=2025-09-19").text

    export = client.get("/records/export.xlsx?form=drug-test")
    ws = load_workbook(io.BytesIO(export.content)).active
    assert ws["B2"].value == "Bob Wilson" and ws["J2"].value == "reviewer"
    assert export.headers["content-disposition"].startswith('attachment; filename="Drugtest-')

    actions = [e["action"] for e in db.audit_entries()]
    for action in ("upload.create", "entry.create", "upload.skip", "records.view", "records.export"):
        assert action in actions


def test_admin_user_management(app, db):
    client = login(app, "admin")
    token = csrf(client, "/admin/users")
    response = client.post("/admin/users", data={"csrf": token, "username": "nurse", "password": "short", "role": "viewer"})
    assert response.status_code == 422
    client.post("/admin/users", data={"csrf": token, "username": "nurse", "password": PASSWORD, "role": "reviewer"})
    assert db.get_user_by_name("nurse")["role"] == "reviewer"

    admin_id = db.get_user_by_name("admin")["id"]
    response = client.post(f"/admin/users/{admin_id}", data={"csrf": token, "role": "viewer"})
    assert response.status_code == 422 and "at least one active admin" in response.text
    nurse_id = db.get_user_by_name("nurse")["id"]
    client.post(f"/admin/users/{nurse_id}", data={"csrf": token, "role": "reviewer", "action": "disable"})
    assert db.get_user(nurse_id)["active"] == 0
    assert "user.disable" in [e["action"] for e in db.audit_entries()]
    assert "user.disable" in client.get("/admin/audit").text


def test_other_form_type_end_to_end(app, db):
    client = login(app, "reviewer")
    token = csrf(client, "/review")
    assert "Equipment inspection" in client.get("/review").text
    client.post("/uploads", data={"csrf": token, "form_key": "equipment-inspection"},
                files=[("files", ("check.png", png(), "image/png"))])
    wait_until_read(db)
    [upload] = db.pending_uploads()
    page = client.get(f"/review/{upload['id']}").text
    assert 'name="field_equipment_id" value="FL-3"' in page and "Needs repair" in page
    form = {"csrf": token, "action": "save", "field_equipment_id": "FL-3", "field_inspection_date": "2025-03-01",
            "field_inspector": "Ana", "field_status": "Needs repair"}
    assert client.post(f"/review/{upload['id']}", data=form, follow_redirects=False).status_code == 303
    page = client.get("/records?form=equipment-inspection&flag=*flagged").text
    assert "FL-3" in page and 'class="flagged"' in page
    assert db.count("drug-test") == 0


def test_upload_rejects_unknown_form(app):
    client = login(app, "reviewer")
    token = csrf(client, "/review")
    response = client.post("/uploads", data={"csrf": token, "form_key": "nope"},
                           files=[("files", ("a.png", png(), "image/png"))])
    assert response.status_code == 400


def test_admin_form_editor(app, db):
    client = login(app, "admin")
    token = csrf(client, "/admin/forms")
    assert login(app, "reviewer").get("/admin/forms").status_code == 403

    response = client.post("/admin/forms", data={"csrf": token, "name": "Vehicle check", "copy_from": ""},
                           follow_redirects=False)
    assert response.headers["location"] == "/admin/forms/vehicle-check"
    assert db.get_form("vehicle-check").keys == ["name"]

    edit = {"csrf": token, "name": "Vehicle check", "description": "Daily pre-use check", "field_count": "1",
            "f0_key": "name", "f0_label": "Driver", "f0_type": "text", "f0_order": "2", "f0_required": "1",
            "f1_label": "Check date", "f1_type": "date", "f1_order": "1", "f1_duplicate": "1",
            "date_field": "", "flag_field": "", "flag_values": ""}
    assert client.post("/admin/forms/vehicle-check", data=edit, follow_redirects=False).status_code == 303
    form = db.get_form("vehicle-check")
    assert [(f.key, f.label) for f in form.fields] == [("check_date", "Check date"), ("name", "Driver")]
    assert form.duplicate_keys == ["check_date"] and form.description == "Daily pre-use check"

    edit.update({"field_count": "2", "f0_key": "check_date", "f0_label": "Check date", "f0_type": "date",
                 "f1_key": "name", "f1_label": "Driver", "f1_type": "text",
                 "f2_label": "Result", "f2_type": "choice", "f2_choices": "OK, Fault",
                 "date_field": "check_date", "flag_field": "result", "flag_values": "fault"})
    client.post("/admin/forms/vehicle-check", data=edit)
    form = db.get_form("vehicle-check")
    assert form.field("result").choices == ["OK", "Fault"] and form.flag_values == ["Fault"]

    bad = {**edit, "f3_label": "Broken", "f3_type": "choice", "f3_choices": ""}
    edit["field_count"] = "3"
    response = client.post("/admin/forms/vehicle-check", data={**bad, "field_count": "3"})
    assert response.status_code == 422 and "needs at least one choice" in response.text

    db.add_entry(form, {"check_date": "2025-01-01", "name": "Al", "result": "OK"}, "test")
    response = client.post("/admin/forms/vehicle-check", data={**edit, "f2_remove": "1"})
    assert response.status_code == 422 and "can't be removed" in response.text
    response = client.post("/admin/forms/vehicle-check/delete", data={"csrf": token})
    assert response.status_code == 422 and db.get_form("vehicle-check")

    client.post("/admin/forms/training-signoff/delete", data={"csrf": token})
    assert db.get_form("training-signoff") is None
    actions = [e["action"] for e in db.audit_entries()]
    assert {"form.create", "form.update", "form.delete"} <= set(actions)
