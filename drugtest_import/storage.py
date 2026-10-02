"""SQLite record store with Excel export."""
import csv
import os
import sqlite3
import tempfile
from datetime import date, datetime, timezone

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill

from .schema import FIELDS, normalize, validate

COLUMNS = {
    "EmployeeID": "employee_id", "Name": "name", "Department": "department", "TestDate": "test_date",
    "TestType": "test_type", "Result": "result", "Notes": "notes",
}

SCHEMA = """
CREATE TABLE IF NOT EXISTS records (
    id INTEGER PRIMARY KEY,
    employee_id TEXT,
    name TEXT NOT NULL,
    department TEXT,
    test_date TEXT NOT NULL,
    test_type TEXT NOT NULL,
    result TEXT NOT NULL,
    notes TEXT,
    source_file TEXT,
    created_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS records_lookup ON records (name, test_date, test_type);
"""

EXPORT_HEADERS = FIELDS + ["SourceFile", "EnteredAt"]
COLUMN_WIDTHS = [12, 24, 16, 12, 10, 13, 40, 24, 20]
RESULT_FILLS = {"Positive": "FECACA", "Pending": "FEF08A", "Inconclusive": "FEF08A", "Refused": "FED7AA"}


class RecordStore:
    def __init__(self, path: str):
        self.path = path
        self.conn = sqlite3.connect(path)
        self.conn.row_factory = sqlite3.Row
        self.conn.executescript(SCHEMA)

    def close(self):
        self.conn.close()

    def add(self, record: dict, source_file: str | None = None) -> int:
        """Insert a validated record. Only the file's base name is stored, not its full path."""
        values = {col: record.get(field) for field, col in COLUMNS.items()}
        values["source_file"] = os.path.basename(source_file) if source_file else None
        values["created_at"] = datetime.now(timezone.utc).isoformat(timespec="seconds")
        with self.conn:
            cursor = self.conn.execute(
                f"INSERT INTO records ({', '.join(values)}) VALUES ({', '.join('?' * len(values))})",
                list(values.values()),
            )
        return cursor.lastrowid

    def find_duplicate(self, record: dict) -> dict | None:
        """Return an existing record for the same person, date and test type, if any."""
        row = self.conn.execute(
            "SELECT * FROM records WHERE lower(name) = lower(?) AND test_date = ? AND test_type = ? "
            "AND (employee_id IS ? OR ? IS NULL OR employee_id IS NULL) LIMIT 1",
            (record.get("Name"), record.get("TestDate"), record.get("TestType"),
             record.get("EmployeeID"), record.get("EmployeeID")),
        ).fetchone()
        return self._to_record(row) if row else None

    def all(self) -> list[dict]:
        rows = self.conn.execute("SELECT * FROM records ORDER BY test_date, id").fetchall()
        return [self._to_record(row) for row in rows]

    def count(self) -> int:
        return self.conn.execute("SELECT count(*) FROM records").fetchone()[0]

    def import_csv(self, path: str) -> tuple[int, int]:
        """Load records from a CSV produced by earlier versions.

        Returns (imported, skipped); rows that fail validation are skipped.
        """
        imported = skipped = 0
        with open(path, newline="", encoding="utf-8-sig") as f:
            for row in csv.DictReader(f):
                record = normalize(row)
                if validate(record):
                    skipped += 1
                    continue
                self.add(record, source_file=path)
                imported += 1
        return imported, skipped

    def export_xlsx(self, path: str) -> int:
        """Write every record to a formatted Excel workbook. Returns rows written."""
        records = self.all()
        wb = Workbook()
        ws = wb.active
        ws.title = "Drug Tests"
        ws.append(EXPORT_HEADERS)
        for record in records:
            row = [record[field] for field in FIELDS] + [record["SourceFile"], record["EnteredAt"]]
            try:
                row[FIELDS.index("TestDate")] = date.fromisoformat(record["TestDate"])
            except (TypeError, ValueError):
                pass
            ws.append(row)

        for cell in ws[1]:
            cell.font = Font(bold=True)
        for index, width in enumerate(COLUMN_WIDTHS, start=1):
            ws.column_dimensions[ws.cell(row=1, column=index).column_letter].width = width
        for cell in ws["D"][1:]:
            cell.number_format = "yyyy-mm-dd"
        # Employee IDs stay text so leading zeros survive
        for cell in ws["A"][1:]:
            cell.number_format = "@"
        ws.freeze_panes = "A2"
        last_row = max(len(records) + 1, 2)
        ws.auto_filter.ref = f"A1:I{last_row}"
        result_col = ws.cell(row=1, column=FIELDS.index("Result") + 1).column_letter
        for result, color in RESULT_FILLS.items():
            ws.conditional_formatting.add(
                f"A2:I{last_row}",
                FormulaRule(formula=[f'${result_col}2="{result}"'], fill=PatternFill("solid", fgColor=color)),
            )

        directory = os.path.dirname(os.path.abspath(path))
        fd, tmp_path = tempfile.mkstemp(dir=directory, suffix=".xlsx")
        os.close(fd)
        try:
            wb.save(tmp_path)
            os.replace(tmp_path, path)
        except BaseException:
            os.unlink(tmp_path)
            raise
        return len(records)

    @staticmethod
    def _to_record(row: sqlite3.Row) -> dict:
        record = {field: row[col] for field, col in COLUMNS.items()}
        record["SourceFile"] = row["source_file"]
        record["EnteredAt"] = row["created_at"]
        return record
